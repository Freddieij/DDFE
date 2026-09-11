"""Reimplementa FuelLayout/TechLayout do .pq sobre RES_Dados_Siglas.xlsx
para conferir o que os nos fantasmas realmente fazem."""
import openpyxl, statistics
from collections import defaultdict

XLSX = r"C:\Users\frede\Desktop\Repositorios\Power_BI_Borba\RES_Dados_Siglas.xlsx"
wb = openpyxl.load_workbook(XLSX)

fuels = []
for r in list(wb["Fuels"].iter_rows(values_only=True))[1:]:
    if r[0]:
        fuels.append({"fuel": str(r[0]).strip(), "stage": int(r[1]), "color": r[2]})
techs = []
for r in list(wb["Techs"].iter_rows(values_only=True))[1:]:
    if any(r[:3]):
        techs.append({
            "tech": (str(r[0]).strip() if r[0] else None),
            "fuel_in": (str(r[1]).strip() if r[1] else None),
            "fuel_out": (str(r[2]).strip() if r[2] else None),
        })

# ---- Input_Tech: logica de hub para linhas sem tech -----------------
for t in techs:
    if not t["tech"]:
        t["tech"] = f"LINK_HUB_{t['fuel_in']}" if t["fuel_in"] else "LINK_ORFAO"

# ---- FuelLayout ----------------------------------------------------
# dedup por fuel, agrupa por stage
seen = set(); fl = []
for f in fuels:
    if f["fuel"] not in seen:
        seen.add(f["fuel"]); fl.append(f)

por_stage = defaultdict(list)
for f in fl:
    por_stage[f["stage"]].append(f)
stages = sorted(por_stage)

GAP, PPI = 400, 30
larguras = []
for i, st in enumerate(stages):
    minha = len(por_stage[st])
    prox = len(por_stage[stages[i + 1]]) if i + 1 < len(stages) else 0
    larguras.append(GAP + max(minha, prox) * PPI)

xbase = {}
acc = 150
for st, w in zip(stages, larguras):
    xbase[st] = acc
    acc += w

x_fuel = {}
for st in stages:
    itens = por_stage[st]
    n = len(itens)
    for idx, f in enumerate(itens, start=1):
        desvio = -0.30 if n == 1 else ((idx / (n + 1)) - 0.5) * 0.9
        x_fuel[f["fuel"]] = round(xbase[st] + desvio * (n * 30 + 150))

stage_of = {f["fuel"]: f["stage"] for f in fl}
# Fuel_Global_ID: ordena por (stage, fuel) e indexa de 1
ordenado = sorted(fl, key=lambda f: (f["stage"], f["fuel"]))
fgid = {f["fuel"]: i for i, f in enumerate(ordenado, start=1)}

parede_dir = {st: max(x_fuel[f["fuel"]] for f in por_stage[st]) for st in stages}
parede_esq = {st: min(x_fuel[f["fuel"]] for f in por_stage[st]) for st in stages}

# ---- TechLayout: agregados por tech --------------------------------
agg = defaultdict(lambda: {"in_st": [], "out_st": [], "in_ids": [], "rows": []})
for t in techs:
    a = agg[t["tech"]]
    a["rows"].append(t)
    if t["fuel_in"] in stage_of:
        a["in_st"].append(stage_of[t["fuel_in"]])
        a["in_ids"].append(fgid[t["fuel_in"]])
    if t["fuel_out"] in stage_of:
        a["out_st"].append(stage_of[t["fuel_out"]])

def m_round(v):  # Number.Round do M = banker's rounding, igual ao round() do Python
    return int(round(v))

info = {}
for tech, a in agg.items():
    max_in = max(a["in_st"]) if a["in_st"] else None
    min_out = min(a["out_st"]) if a["out_st"] else None
    raia = m_round(statistics.fmean(a["in_ids"])) if a["in_ids"] else 999
    pd_in = parede_dir.get(max_in) if max_in else None
    pe_out = parede_esq.get(min_out) if min_out else None
    pd_prev = parede_dir.get(min_out - 1) if min_out else None
    pe_next = parede_esq.get(max_in + 1) if max_in else None

    if max_in is not None and min_out is not None:
        if max_in >= min_out:
            x = (pd_prev + pe_out) / 2 if pd_prev is not None else pe_out - 200
        else:
            x = (pd_in + pe_next) / 2 if pe_next is not None else pd_in + 150
    elif max_in is not None:
        x = (pd_in + pe_next) / 2 if pe_next is not None else pd_in + 150
    elif min_out is not None:
        x = (pd_prev + pe_out) / 2 if pd_prev is not None else pe_out - 150
    else:
        x = 0
    info[tech] = {"x_tech": round(x), "raia": raia, "x_grid": m_round(round(x) / 50) * 50,
                  "backflow": (max_in is not None and min_out is not None and max_in >= min_out)}

# ---- Nos fantasmas -------------------------------------------------
grids_ocupados = sorted({info[t]["x_grid"] for t in info})

def monta(com_fantasmas):
    linhas = [{"tech": t, "raia": info[t]["raia"], "x_grid": info[t]["x_grid"], "ghost": False}
              for t in info]
    if com_fantasmas:
        n = 0
        for t in techs:
            tech = t["tech"]
            xo = x_fuel.get(t["fuel_out"])
            xt = info[tech]["x_tech"]
            if xo is None or abs(xo - xt) <= 70:
                continue
            lo, hi = min(xt, xo), max(xt, xo)
            for g in grids_ocupados:
                if lo + 25 < g < hi - 25:
                    n += 1
                    linhas.append({"tech": f"{tech}_GHOST_{n}", "raia": info[tech]["raia"],
                                   "x_grid": g, "ghost": True})
    # densidade por (raia, x_grid) contando techs distintos
    dens = defaultdict(set)
    for l in linhas:
        dens[(l["raia"], l["x_grid"])].add(l["tech"])
    qtd = {k: len(v) for k, v in dens.items()}
    altura = defaultdict(int)
    for (raia, _), q in qtd.items():
        altura[raia] = max(altura[raia], q)
    ybase = {}
    acc = 1
    for raia in sorted(altura):
        ybase[raia] = acc
        acc += altura[raia]
    # indice interno: reais primeiro (ghost False < True), depois nome
    y = {}
    grupos = defaultdict(list)
    for l in linhas:
        grupos[(l["x_grid"], l["raia"])].append(l)
    for (xg, raia), itens in grupos.items():
        vistos = {}
        for l in itens:
            vistos.setdefault(l["tech"], l)
        ordem = sorted(vistos.values(), key=lambda l: (l["ghost"], l["tech"]))
        for i, l in enumerate(ordem):
            y[l["tech"]] = ybase[raia] + i
    return y, altura, ybase

for rotulo, com in (("SEM fantasmas", False), ("COM fantasmas", True)):
    y, altura, ybase = monta(com)
    # colisoes: conexao longa de 'tech' passa por caixa real 'outro' na mesma linha y
    colisoes = []
    for t in techs:
        tech = t["tech"]
        xo = x_fuel.get(t["fuel_out"])
        xt = info[tech]["x_tech"]
        if xo is None:
            continue
        lo, hi = min(xt, xo), max(xt, xo)
        yt = y.get(tech)
        for outro in info:
            if outro == tech:
                continue
            xg = info[outro]["x_grid"]
            if lo + 25 < xg < hi - 25 and y.get(outro) == yt:
                colisoes.append((tech, outro, yt))
    print(f"{rotulo}: altura total={sum(altura.values())}  "
          f"colisoes linha-x-caixa={len(set(colisoes))}")
    for c in sorted(set(colisoes))[:8]:
        print("    linha de", c[0], "atravessa", c[1], "na linha y =", c[2])
